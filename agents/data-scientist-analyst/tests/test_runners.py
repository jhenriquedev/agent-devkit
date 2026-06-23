#!/usr/bin/env python3
"""Tests for the data-scientist-analyst agent runners."""

from __future__ import annotations

import json
import importlib.util
import inspect
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace


ROOT = Path(__file__).resolve().parents[3]
CLI = ROOT / "ai-devkit"


class DataScientistAnalystRunnersTest(unittest.TestCase):
    def run_cli(self, *args: str) -> subprocess.CompletedProcess[str]:
        return subprocess.run(
            [sys.executable, str(CLI), *args],
            cwd=ROOT,
            check=False,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )

    def parse_payload(self, result: subprocess.CompletedProcess[str]) -> dict:
        payload = json.loads(result.stdout)
        if "stdout" in payload:
            return json.loads(payload["stdout"])
        return payload

    def load_runner_support_module(self):
        module_path = ROOT / "agents/data-scientist-analyst/capabilities/_shared/runner_support.py"
        spec = importlib.util.spec_from_file_location("data_scientist_runner_support", module_path)
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module

    def load_dataset_module(self, module_name: str):
        module_path = ROOT / f"agents/data-scientist-analyst/infra/integrations/file-dataset/{module_name}.py"
        module_dir = str(module_path.parent)
        if module_dir not in sys.path:
            sys.path.insert(0, module_dir)
        spec = importlib.util.spec_from_file_location(module_name, module_path)
        self.assertIsNotNone(spec, module_name)
        self.assertIsNotNone(spec.loader, module_name)
        module = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = module
        spec.loader.exec_module(module)
        return module

    def write_customer_csv(self, root: Path) -> Path:
        source = root / "customers.csv"
        source.write_text(
            "\n".join(
                [
                    "cpf,name,status,amount,created_at,email",
                    "12345678909,Ana,active,100.50,2026-06-01,ana@example.com",
                    "11144477735,Beto,inactive,200.00,2026-06-02,beto@example.com",
                    "12345678909,Ana,active,100.50,2026-06-01,ana@example.com",
                    "00000000000,Caio,active,,2026-06-03,caio@example.com",
                ]
            ),
            encoding="utf-8",
        )
        return source

    def write_analytics_csv(self, root: Path) -> Path:
        source = root / "analytics.csv"
        source.write_text(
            "\n".join(
                [
                    "id,segment,age,score,income,converted",
                    "1,A,20,40,100,0",
                    "2,A,30,60,150,1",
                    "3,B,40,80,200,1",
                    "4,B,50,100,250,1",
                    "5,C,60,120,10000,0",
                ]
            ),
            encoding="utf-8",
        )
        return source

    def write_time_series_csv(self, root: Path) -> Path:
        source = root / "events.csv"
        source.write_text(
            "\n".join(
                [
                    "customer_id,signup_date,event_date,amount,converted",
                    "1,2026-01-01,2026-01-01,10,1",
                    "1,2026-01-01,2026-01-02,12,0",
                    "2,2026-01-02,2026-01-02,14,1",
                    "2,2026-01-02,2026-01-03,16,1",
                    "3,2026-01-03,2026-01-03,100,0",
                    "3,2026-01-03,2026-01-04,18,1",
                    "4,2026-01-04,2026-01-04,20,0",
                ]
            ),
            encoding="utf-8",
        )
        return source

    def write_experiment_csv(self, root: Path) -> Path:
        source = root / "experiment.csv"
        source.write_text(
            "\n".join(
                [
                    "user_id,variant,revenue,converted",
                    "1,control,10,0",
                    "2,control,12,0",
                    "3,control,14,1",
                    "4,control,16,1",
                    "5,treatment,18,1",
                    "6,treatment,20,1",
                    "7,treatment,22,1",
                    "8,treatment,24,1",
                ]
            ),
            encoding="utf-8",
        )
        return source

    def write_modeling_csv(self, root: Path) -> Path:
        source = root / "modeling.csv"
        source.write_text(
            "\n".join(
                [
                    "user_id,score,income,channel,converted,converted_copy,post_event_status",
                    "1,10,100,email,0,0,rejected",
                    "2,20,120,email,0,0,rejected",
                    "3,30,140,branch,0,0,rejected",
                    "4,40,300,branch,1,1,approved",
                    "5,50,320,app,1,1,approved",
                    "6,60,340,app,1,1,approved",
                ]
            ),
            encoding="utf-8",
        )
        return source

    def test_data_repository_is_facade_over_domain_modules(self) -> None:
        expected_exports = {
            "dataset_models": ["Dataset", "DataScientistError"],
            "dataset_io": ["read_csv", "read_json", "read_jsonl", "read_xlsx", "hash_file"],
            "profiling": ["profile_column", "infer_type", "quality_score"],
            "privacy": ["detect_sensitive", "mask_value", "mask_row"],
            "statistics_tools": ["summarize_metric", "calculate_cohens_d", "explain_significance"],
            "time_series_tools": ["aggregate_time_series", "parse_date_value", "time_series_data_quality"],
            "modeling_tools": ["train_numeric_threshold_model", "classification_metrics", "detect_leakage_candidates"],
            "reconciliation_tools": ["index_rows", "compare_rows", "normalize_key_value"],
            "reporting": ["render_data_report", "render_reconciliation_report"],
        }
        for module_name, exports in expected_exports.items():
            module = self.load_dataset_module(module_name)
            for export in exports:
                self.assertTrue(hasattr(module, export), f"{module_name}.{export} missing")

        repository = self.load_dataset_module("data_repository")
        repository_functions = {
            name
            for name, value in vars(repository).items()
            if inspect.isfunction(value) and getattr(value, "__module__", "") == "data_repository"
        }
        self.assertFalse(
            repository_functions
            & {
                "read_csv",
                "profile_column",
                "detect_sensitive",
                "aggregate_time_series",
                "train_numeric_threshold_model",
                "index_rows",
                "render_data_report",
            }
        )

    def test_agent_is_listed_with_mvp_capabilities(self) -> None:
        agents_result = self.run_cli("--json", "agents")
        self.assertEqual(agents_result.returncode, 0, agents_result.stderr)
        agents = {item["id"] for item in json.loads(agents_result.stdout)["items"]}
        self.assertIn("data-scientist-analyst", agents)

        capabilities_result = self.run_cli(
            "--json",
            "capabilities",
            "data-scientist-analyst",
        )
        self.assertEqual(capabilities_result.returncode, 0, capabilities_result.stderr)
        capabilities = {
            item["id"].split(".")[-1]
            for item in json.loads(capabilities_result.stdout)["items"]
        }
        self.assertGreaterEqual(
            capabilities,
            {
                "ingest-dataset",
                "inspect-dataset-schema",
                "profile-dataset",
                "detect-sensitive-data",
                "reconcile-spreadsheets",
                "generate-reconciliation-report",
                "generate-data-report",
                "analyze-sql-source",
                "run-data-pipeline",
                "run-exploratory-analysis",
                "detect-outliers",
                "analyze-correlation",
                "segment-data",
                "analyze-time-series",
                "compare-periods",
                "analyze-cohorts",
                "detect-anomalies",
                "forecast-series",
                "test-hypothesis",
                "calculate-confidence-intervals",
                "calculate-sample-size",
                "measure-effect-size",
                "explain-statistical-result",
                "prepare-modeling-dataset",
                "baseline-predictive-model",
                "evaluate-model",
                "explain-model-results",
                "detect-data-leakage",
                "monitor-model-drift",
            },
        )

    def test_operational_docs_and_contracts_are_versioned(self) -> None:
        agent_root = ROOT / "agents" / "data-scientist-analyst"
        readme = (agent_root / "README.md").read_text(encoding="utf-8")
        runbook = (agent_root / "knowledge" / "operational-runbook.md").read_text(encoding="utf-8")
        health = (agent_root / "knowledge" / "health-checklist.md").read_text(encoding="utf-8")

        self.assertIn("## Operacao", readme)
        self.assertIn("--max-rows", readme)
        self.assertIn("run-data-pipeline", readme)
        self.assertIn("Quality gates", runbook)
        self.assertIn("Troubleshooting", runbook)
        self.assertIn("Riscos residuais", health)

        contract_dir = agent_root / "knowledge" / "contracts"
        for name in [
            "profile-dataset.schema.json",
            "run-data-pipeline.schema.json",
            "evaluate-model.schema.json",
            "analyze-sql-source.schema.json",
        ]:
            contract = json.loads((contract_dir / name).read_text(encoding="utf-8"))
            self.assertEqual(contract["schema_version"], "0.1.0")
            self.assertIn("required_top_level_keys", contract)
            self.assertTrue(contract["quality_gates"])
        sql_contract = json.loads((contract_dir / "analyze-sql-source.schema.json").read_text(encoding="utf-8"))
        self.assertGreaterEqual(
            set(sql_contract["required_top_level_keys"]),
            {"delegation", "analysis", "quality_gates"},
        )

    def test_profile_dataset_reports_schema_quality_and_sensitive_fields(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_customer_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "profile-dataset",
                "--source",
                str(source),
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["dataset"]["row_count"], 4)
        self.assertEqual(payload["dataset"]["column_count"], 6)
        self.assertEqual(payload["quality"]["duplicate_row_count"], 1)
        self.assertIn("cpf", payload["schema"]["probable_keys"])
        self.assertIn("cpf", payload["sensitive_data"]["columns"])
        self.assertEqual(payload["columns"]["amount"]["missing_count"], 1)
        self.assertEqual(payload["columns"]["amount"]["inferred_type"], "number")

    def test_reconcile_spreadsheets_matches_and_explains_differences(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            left = root / "erp.csv"
            right = root / "bank.csv"
            left.write_text(
                "\n".join(
                    [
                        "id,cpf,amount,status",
                        "1,12345678909,100.00,paid",
                        "2,11144477735,200.00,pending",
                        "3,55566677788,50.00,paid",
                    ]
                ),
                encoding="utf-8",
            )
            right.write_text(
                "\n".join(
                    [
                        "id,cpf,amount,status",
                        "1,12345678909,100.01,paid",
                        "2,11144477735,250.00,paid",
                        "4,99988877766,10.00,paid",
                    ]
                ),
                encoding="utf-8",
            )
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "reconcile-spreadsheets",
                "--left",
                str(left),
                "--right",
                str(right),
                "--key",
                "id",
                "--compare-columns",
                "amount,status",
                "--numeric-tolerance",
                "0.02",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["summary"]["matched_count"], 1)
        self.assertEqual(payload["summary"]["mismatched_count"], 1)
        self.assertEqual(payload["summary"]["missing_right_count"], 1)
        self.assertEqual(payload["summary"]["missing_left_count"], 1)
        mismatch = payload["mismatched"][0]
        self.assertEqual(mismatch["key"], "2")
        fields = {item["column"] for item in mismatch["differences"]}
        self.assertEqual(fields, {"amount", "status"})

    def test_generate_data_report_writes_markdown_artifact(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = self.write_customer_csv(root)
            output = root / "report.md"
            result = self.run_cli(
                "run",
                "data-scientist-analyst",
                "generate-data-report",
                "--source",
                str(source),
                "--output",
                str(output),
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            content = output.read_text(encoding="utf-8")

        self.assertIn("# Relatorio de Dados", content)
        self.assertIn("Linhas: 4", content)
        self.assertIn("Dados sensiveis", content)
        self.assertIn("## Reprodutibilidade", content)
        self.assertIn("## Limitacoes", content)
        self.assertIn("## Base para PDF", content)

    def test_common_json_output_creates_parent_directories(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = self.write_customer_csv(root)
            output = root / "nested" / "results" / "profile.json"
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "profile-dataset",
                "--source",
                str(source),
                "--output",
                str(output),
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue(output.exists())
            written = json.loads(output.read_text(encoding="utf-8"))

        self.assertEqual(written["dataset"]["row_count"], 4)

    def test_profile_dataset_reads_selected_xlsx_sheet(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "multi-sheet.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.title = "Ignored"
            first.append(["id", "amount"])
            first.append([1, 999])
            selected = workbook.create_sheet("Selected")
            selected.append(["id", "amount"])
            selected.append([10, 100])
            selected.append([20, 200])
            workbook.save(source)
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "profile-dataset",
                "--source",
                str(source),
                "--sheet",
                "Selected",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["dataset"]["row_count"], 2)
        self.assertEqual(payload["columns"]["amount"]["numeric"]["mean"], 150.0)

    def test_profile_dataset_reads_json_path_records(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "nested.json"
            source.write_text(
                json.dumps(
                    {
                        "metadata": {"source": "unit-test"},
                        "payload": {
                            "items": [
                                {"id": 1, "amount": 10},
                                {"id": 2, "amount": 20},
                            ]
                        },
                    }
                ),
                encoding="utf-8",
            )
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "profile-dataset",
                "--source",
                str(source),
                "--json-path",
                "payload.items",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["dataset"]["row_count"], 2)
        self.assertEqual(payload["columns"]["amount"]["numeric"]["mean"], 15.0)

    def test_detect_sensitive_data_does_not_classify_phone_as_cpf(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = Path(tmpdir) / "contacts.csv"
            source.write_text(
                "id,email,phone\n1,a@example.com,11999998888\n",
                encoding="utf-8",
            )
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "detect-sensitive-data",
                "--source",
                str(source),
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["sensitive_data"]["columns"]["phone"], ["phone"])

    def test_detect_sensitive_data_masks_names_and_preserves_phone_kind(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "people.csv"
            source.write_text(
                "\n".join(
                    [
                        "nome,phone,customer_id",
                        "Ana Maria,+5511999998888,12345678901",
                        "Bruno Lima,+5511888887777,abc-123",
                    ]
                ),
                encoding="utf-8",
            )
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "detect-sensitive-data",
                "--source",
                str(source),
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["sensitive_data"]["columns"]["nome"], ["person_name"])
        self.assertEqual(payload["sensitive_data"]["masked_examples"]["nome"][0], "A***")
        self.assertEqual(payload["sensitive_data"]["columns"]["phone"], ["phone"])
        self.assertNotIn("customer_id", payload["sensitive_data"]["columns"])

    def test_profile_dataset_applies_max_rows_and_reports_truncation(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "large.csv"
            source.write_text(
                "\n".join(["id,value", *[f"{index},{index * 10}" for index in range(1, 8)]]),
                encoding="utf-8",
            )
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "profile-dataset",
                "--source",
                str(source),
                "--max-rows",
                "3",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["dataset"]["row_count"], 3)
        self.assertTrue(payload["dataset"]["truncated"])
        self.assertEqual(payload["dataset"]["original_row_count"], 7)
        self.assertTrue(payload["dataset"]["warnings"])

    def test_detect_outliers_reports_iqr_outlier(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_analytics_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "detect-outliers",
                "--source",
                str(source),
                "--columns",
                "income",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        income = payload["columns"]["income"]
        self.assertEqual(income["outlier_count"], 1)
        self.assertEqual(income["outliers"][0]["row_number"], 5)
        self.assertEqual(income["outliers"][0]["value"], 10000.0)

    def test_analyze_correlation_ranks_numeric_relationships(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_analytics_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "analyze-correlation",
                "--source",
                str(source),
                "--columns",
                "age,score,income",
                "--target-column",
                "score",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        top_pair = payload["top_correlations"][0]
        self.assertEqual({top_pair["left"], top_pair["right"]}, {"age", "score"})
        self.assertAlmostEqual(top_pair["correlation"], 1.0)
        target_pairs = {
            item["column"]: item["correlation"]
            for item in payload["target_correlations"]
        }
        self.assertAlmostEqual(target_pairs["age"], 1.0)

    def test_segment_data_summarizes_category_metrics(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_analytics_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "segment-data",
                "--source",
                str(source),
                "--segment-column",
                "segment",
                "--metric-column",
                "converted",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        segments = {item["segment"]: item for item in payload["segments"]}
        self.assertEqual(segments["A"]["row_count"], 2)
        self.assertEqual(segments["A"]["metric"]["sum"], 1.0)
        self.assertEqual(segments["B"]["metric"]["mean"], 1.0)
        self.assertEqual(payload["summary"]["segment_count"], 3)

    def test_run_exploratory_analysis_combines_phase_one_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_analytics_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "run-exploratory-analysis",
                "--source",
                str(source),
                "--target-column",
                "converted",
                "--segment-column",
                "segment",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["dataset"]["row_count"], 5)
        self.assertIn("income", payload["outliers"]["columns"])
        self.assertTrue(payload["correlations"]["top_correlations"])
        self.assertEqual(payload["segments"]["summary"]["segment_count"], 3)
        self.assertTrue(payload["hypotheses"])

    def test_analyze_time_series_aggregates_metric_by_day(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_time_series_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "analyze-time-series",
                "--source",
                str(source),
                "--date-column",
                "event_date",
                "--metric-column",
                "amount",
                "--granularity",
                "day",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["summary"]["period_count"], 4)
        self.assertEqual(payload["series"][0]["period"], "2026-01-01")
        self.assertEqual(payload["series"][0]["sum"], 10.0)
        self.assertEqual(payload["series"][2]["sum"], 116.0)
        self.assertEqual(payload["trend"]["direction"], "up")

    def test_analyze_time_series_reports_invalid_date_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "events.csv"
            source.write_text(
                "\n".join(
                    [
                        "event_date,amount",
                        "2026-01-01,10",
                        "01/02/2026,20",
                        "2026-03-01T10:20:30-03:00,30",
                        "not-a-date,40",
                    ]
                ),
                encoding="utf-8",
            )
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "analyze-time-series",
                "--source",
                str(source),
                "--date-column",
                "event_date",
                "--metric-column",
                "amount",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["summary"]["period_count"], 3)
        self.assertEqual(payload["data_quality"]["invalid_date_rows"], 1)
        self.assertTrue(payload["warnings"])

    def test_compare_periods_reports_delta_between_date_ranges(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_time_series_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "compare-periods",
                "--source",
                str(source),
                "--date-column",
                "event_date",
                "--metric-column",
                "amount",
                "--baseline-start",
                "2026-01-01",
                "--baseline-end",
                "2026-01-02",
                "--comparison-start",
                "2026-01-03",
                "--comparison-end",
                "2026-01-04",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["baseline"]["sum"], 36.0)
        self.assertEqual(payload["comparison"]["sum"], 154.0)
        self.assertEqual(payload["delta"]["absolute"], 118.0)
        self.assertAlmostEqual(payload["delta"]["percent"], 327.777778)

    def test_analyze_cohorts_groups_by_signup_age(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_time_series_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "analyze-cohorts",
                "--source",
                str(source),
                "--cohort-column",
                "signup_date",
                "--event-date-column",
                "event_date",
                "--metric-column",
                "converted",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["summary"]["cohort_count"], 4)
        first = payload["cohorts"][0]
        self.assertEqual(first["cohort"], "2026-01-01")
        self.assertEqual(first["periods"]["0"]["row_count"], 1)
        self.assertEqual(first["periods"]["0"]["metric"]["sum"], 1.0)
        self.assertEqual(first["periods"]["1"]["metric"]["sum"], 0.0)

    def test_detect_anomalies_flags_spike_in_time_series(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_time_series_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "detect-anomalies",
                "--source",
                str(source),
                "--date-column",
                "event_date",
                "--metric-column",
                "amount",
                "--threshold",
                "1.0",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        periods = [item["period"] for item in payload["anomalies"]]
        self.assertIn("2026-01-03", periods)

    def test_detect_anomalies_reports_quality_and_short_series_warning(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "short_series.csv"
            source.write_text(
                "\n".join(
                    [
                        "event_date,amount",
                        "2026-01-01,10",
                        "2026-01-02,11",
                        "invalid,12",
                    ]
                ),
                encoding="utf-8",
            )
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "detect-anomalies",
                "--source",
                str(source),
                "--date-column",
                "event_date",
                "--metric-column",
                "amount",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["data_quality"]["invalid_date_rows"], 1)
        self.assertTrue(any(item["code"] == "short_time_series" for item in payload["validity_warnings"]))

    def test_forecast_series_projects_next_periods(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_time_series_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "forecast-series",
                "--source",
                str(source),
                "--date-column",
                "event_date",
                "--metric-column",
                "amount",
                "--periods",
                "2",
                "--window",
                "2",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(len(payload["forecast"]), 2)
        self.assertEqual(payload["forecast"][0]["period"], "2026-01-05")
        self.assertEqual(payload["forecast"][0]["forecast"], 77.0)
        self.assertEqual(payload["method"], "moving_average")

    def test_test_hypothesis_compares_group_means(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_experiment_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "test-hypothesis",
                "--source",
                str(source),
                "--test-type",
                "mean-difference",
                "--group-column",
                "variant",
                "--group-a",
                "control",
                "--group-b",
                "treatment",
                "--metric-column",
                "revenue",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["test_type"], "mean-difference")
        self.assertEqual(payload["groups"]["control"]["mean"], 13.0)
        self.assertEqual(payload["groups"]["treatment"]["mean"], 21.0)
        self.assertEqual(payload["difference"], 8.0)
        self.assertTrue(payload["p_value"] < 0.01)
        self.assertTrue(payload["significant"])

    def test_test_hypothesis_reports_assumptions_and_missing_values(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "experiment_missing.csv"
            source.write_text(
                "\n".join(
                    [
                        "variant,revenue",
                        "control,10",
                        "control,",
                        "control,14",
                        "treatment,18",
                        "treatment,20",
                        "treatment,",
                    ]
                ),
                encoding="utf-8",
            )
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "test-hypothesis",
                "--source",
                str(source),
                "--test-type",
                "mean-difference",
                "--group-column",
                "variant",
                "--group-a",
                "control",
                "--group-b",
                "treatment",
                "--metric-column",
                "revenue",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["data_quality"]["missing_metric_rows"], 2)
        self.assertTrue(payload["assumptions"])
        self.assertTrue(payload["validity_warnings"])
        self.assertEqual(payload["recommended_next_test"], "welch_t_test_or_nonparametric_test")

    def test_calculate_confidence_intervals_reports_mean_interval(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_experiment_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "calculate-confidence-intervals",
                "--source",
                str(source),
                "--metric-column",
                "revenue",
                "--confidence",
                "0.95",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["metric_column"], "revenue")
        self.assertEqual(payload["mean"], 17.0)
        self.assertEqual(payload["sample_size"], 8)
        self.assertAlmostEqual(payload["interval"]["lower"], 13.605, places=3)
        self.assertAlmostEqual(payload["interval"]["upper"], 20.395, places=3)

    def test_calculate_confidence_intervals_warns_for_small_sample_and_missing_values(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "small.csv"
            source.write_text("revenue\n10\n \n12\n14\n", encoding="utf-8")
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "calculate-confidence-intervals",
                "--source",
                str(source),
                "--metric-column",
                "revenue",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["data_quality"]["missing_metric_rows"], 1)
        self.assertTrue(any(item["code"] == "small_sample" for item in payload["validity_warnings"]))
        self.assertTrue(payload["assumptions"])

    def test_calculate_sample_size_for_two_proportions(self) -> None:
        result = self.run_cli(
            "--json",
            "run",
            "data-scientist-analyst",
            "calculate-sample-size",
            "--baseline-rate",
            "0.5",
            "--minimum-detectable-effect",
            "0.1",
            "--alpha",
            "0.05",
            "--power",
            "0.8",
        )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["method"], "two_proportions_normal_approx")
        self.assertEqual(payload["sample_size_per_group"], 393)
        self.assertEqual(payload["total_sample_size"], 786)

    def test_measure_effect_size_reports_cohens_d(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_experiment_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "measure-effect-size",
                "--source",
                str(source),
                "--group-column",
                "variant",
                "--group-a",
                "control",
                "--group-b",
                "treatment",
                "--metric-column",
                "revenue",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["method"], "cohens_d")
        self.assertAlmostEqual(payload["effect_size"], 3.098387, places=6)
        self.assertEqual(payload["magnitude"], "large")

    def test_explain_statistical_result_interprets_significance(self) -> None:
        result = self.run_cli(
            "--json",
            "run",
            "data-scientist-analyst",
            "explain-statistical-result",
            "--p-value",
            "0.01",
            "--alpha",
            "0.05",
            "--effect-size",
            "0.8",
        )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertTrue(payload["significant"])
        self.assertEqual(payload["effect_magnitude"], "large")
        self.assertIn("estatisticamente significativo", payload["executive_summary"])

    def test_prepare_modeling_dataset_selects_features_and_splits(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_modeling_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "prepare-modeling-dataset",
                "--source",
                str(source),
                "--target-column",
                "converted",
                "--feature-columns",
                "score,income,channel",
                "--test-size",
                "0.33",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["target"]["column"], "converted")
        self.assertEqual(payload["target"]["type"], "classification")
        self.assertEqual(payload["features"]["selected"], ["score", "income", "channel"])
        self.assertEqual(payload["split"]["train_rows"], 4)
        self.assertEqual(payload["split"]["test_rows"], 2)

    def test_baseline_predictive_model_finds_numeric_threshold(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_modeling_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "baseline-predictive-model",
                "--source",
                str(source),
                "--target-column",
                "converted",
                "--feature-columns",
                "score,income",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["model"]["type"], "numeric_threshold")
        self.assertEqual(payload["model"]["feature"], "score")
        self.assertEqual(payload["evaluation_scope"], "holdout_baseline")
        self.assertIn("train_metrics", payload)
        self.assertIn("test_metrics", payload)
        self.assertEqual(payload["metrics"], payload["test_metrics"])

    def test_evaluate_model_reports_classification_metrics(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_modeling_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "evaluate-model",
                "--source",
                str(source),
                "--target-column",
                "converted",
                "--feature-columns",
                "score,income",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["evaluation_scope"], "holdout_baseline")
        self.assertEqual(payload["metrics"], payload["test_metrics"])
        self.assertIn("train_metrics", payload)
        self.assertIn("confusion_matrix", payload["test_metrics"])
        self.assertIn("quality_gates", payload)
        self.assertTrue(any(gate["name"] == "holdout_available" for gate in payload["quality_gates"]))

    def test_evaluate_model_reports_class_balance_and_configurable_holdout(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "imbalanced_modeling.csv"
            source.write_text(
                "\n".join(
                    [
                        "score,converted",
                        "10,0",
                        "20,0",
                        "30,0",
                        "40,0",
                        "50,0",
                        "60,0",
                        "70,1",
                        "80,0",
                        "90,1",
                        "100,0",
                    ]
                ),
                encoding="utf-8",
            )
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "evaluate-model",
                "--source",
                str(source),
                "--target-column",
                "converted",
                "--feature-columns",
                "score",
                "--test-size",
                "0.3",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["split"]["test_size"], 0.3)
        self.assertEqual(payload["class_balance"]["total"]["0"], 8)
        self.assertEqual(payload["class_balance"]["total"]["1"], 2)
        self.assertTrue(any(item["code"] == "imbalanced_classes" for item in payload["validity_warnings"]))
        self.assertIn("balanced_accuracy", payload["test_metrics"])

    def test_analyze_sql_source_normalizes_tabular_rows(self) -> None:
        module = self.load_runner_support_module()

        result = module.normalize_sql_result(
            {
                "table": "customers",
                "columns": ["id", "amount"],
                "rows": [{"id": 1, "amount": 10.5}, {"id": 2, "amount": 20}],
            }
        )

        self.assertEqual(result["kind"], "tabular_dataset")
        self.assertEqual(result["dataset"]["row_count"], 2)
        self.assertEqual(result["columns"], ["id", "amount"])
        self.assertEqual(result["rows"][0]["amount"], "10.5")

    def test_sql_tabular_result_writes_reusable_dataset_artifact(self) -> None:
        module = self.load_dataset_module("sql_result")

        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "artifacts" / "sql-dataset.json"
            normalized = module.normalize_sql_result(
                {
                    "table": "customers",
                    "columns": ["id", "amount"],
                    "rows": [{"id": 1, "amount": 10.5}, {"id": 2, "amount": 20}],
                }
            )
            artifact = module.write_tabular_artifact(normalized, str(output))

            self.assertTrue(output.exists())
            payload = json.loads(output.read_text(encoding="utf-8"))

        self.assertEqual(artifact["format"], "json")
        self.assertEqual(payload["rows"][0]["amount"], "10.5")

    def test_analyze_sql_source_returns_operational_contract(self) -> None:
        module = self.load_runner_support_module()

        class Completed:
            returncode = 0
            stderr = ""
            stdout = json.dumps(
                {
                    "stdout": json.dumps(
                        {
                            "table": "customers",
                            "columns": ["id", "amount"],
                            "rows": [{"id": 1, "amount": 10.5}],
                        }
                    )
                }
            )

        original_run = module.subprocess.run
        module.subprocess.run = lambda *args, **kwargs: Completed()
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                args = SimpleNamespace(
                    database_agent="postgres-data-analyzer",
                    database_capability="query-table",
                    database=None,
                    schema=None,
                    table="customers",
                    query=None,
                    limit=None,
                    dataset_output=str(Path(tmpdir) / "sql-dataset.json"),
                )
                payload = module.analyze_sql_source(args)
        finally:
            module.subprocess.run = original_run

        self.assertEqual(payload["delegation"]["status"], "success")
        self.assertEqual(payload["analysis"]["kind"], "tabular_dataset")
        self.assertTrue(payload["quality_gates"])
        self.assertEqual(payload["artifact"]["format"], "json")

    def test_run_data_pipeline_generates_profile_eda_and_report_artifacts(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = self.write_analytics_csv(root)
            output_dir = root / "pipeline"
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "run-data-pipeline",
                "--source",
                str(source),
                "--target-column",
                "converted",
                "--segment-column",
                "segment",
                "--output",
                str(output_dir),
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            payload = self.parse_payload(result)
            manifest_path = Path(payload["artifacts"]["manifest"])
            profile_path = Path(payload["artifacts"]["profile"])
            exploratory_path = Path(payload["artifacts"]["exploratory"])
            report_path = Path(payload["artifacts"]["report"])

            self.assertTrue(manifest_path.exists())
            self.assertTrue(profile_path.exists())
            self.assertTrue(exploratory_path.exists())
            self.assertTrue(report_path.exists())
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))

        self.assertEqual(payload["pipeline"]["steps"], ["ingest", "profile", "exploratory_analysis", "data_report"])
        self.assertIn(payload["pipeline"]["status"], {"success", "warning"})
        self.assertIn("created_at", payload["pipeline"])
        self.assertIn("inputs", payload)
        self.assertIn("quality_gates", payload)
        self.assertEqual(manifest["dataset"]["row_count"], 5)
        self.assertEqual(manifest["artifacts"]["profile"], str(profile_path))
        self.assertEqual(manifest["quality_gates"], payload["quality_gates"])

    def test_explain_model_results_describes_selected_driver(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_modeling_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "explain-model-results",
                "--source",
                str(source),
                "--target-column",
                "converted",
                "--feature-columns",
                "score,income",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        self.assertEqual(payload["primary_driver"]["feature"], "score")
        self.assertIn("score", payload["executive_summary"])
        self.assertTrue(payload["limitations"])

    def test_detect_data_leakage_flags_target_copy_and_post_event_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_modeling_csv(Path(tmpdir))
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "detect-data-leakage",
                "--source",
                str(source),
                "--target-column",
                "converted",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        columns = {item["column"] for item in payload["leakage_candidates"]}
        self.assertIn("converted_copy", columns)
        self.assertIn("post_event_status", columns)

    def test_monitor_model_drift_flags_shifted_numeric_distribution(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            reference = root / "reference.csv"
            current = root / "current.csv"
            reference.write_text("score,income\n10,100\n20,120\n30,140\n", encoding="utf-8")
            current.write_text("score,income\n100,400\n110,420\n120,440\n", encoding="utf-8")
            result = self.run_cli(
                "--json",
                "run",
                "data-scientist-analyst",
                "monitor-model-drift",
                "--reference-source",
                str(reference),
                "--source",
                str(current),
                "--columns",
                "score,income",
            )

        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        drifted = {item["column"] for item in payload["drifted_columns"]}
        self.assertIn("score", drifted)
        self.assertIn("income", drifted)

    # P1.5 — schema contract validation tests
    def _load_schema(self, name: str) -> dict:
        contract_dir = ROOT / "agents" / "data-scientist-analyst" / "knowledge" / "contracts"
        return json.loads((contract_dir / name).read_text(encoding="utf-8"))

    def test_profile_dataset_output_satisfies_contract(self) -> None:
        schema = self._load_schema("profile-dataset.schema.json")
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_customer_csv(Path(tmpdir))
            result = self.run_cli(
                "--json", "run", "data-scientist-analyst", "profile-dataset",
                "--source", str(source),
            )
        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        for key in schema["required_top_level_keys"]:
            self.assertIn(key, payload, f"profile-dataset: missing top-level key '{key}'")
        for key in schema["required_dataset_keys"]:
            self.assertIn(key, payload["dataset"], f"profile-dataset: missing dataset key '{key}'")
        self.assertTrue(payload["dataset"]["sha256"], "sha256 must be non-empty")
        self.assertIsNotNone(payload["quality"].get("quality_score"), "quality_score must be present")

    def test_evaluate_model_output_satisfies_contract(self) -> None:
        schema = self._load_schema("evaluate-model.schema.json")
        with tempfile.TemporaryDirectory() as tmpdir:
            source = self.write_modeling_csv(Path(tmpdir))
            result = self.run_cli(
                "--json", "run", "data-scientist-analyst", "evaluate-model",
                "--source", str(source),
                "--target-column", "converted",
                "--feature-columns", "score,income",
            )
        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        for key in schema["required_top_level_keys"]:
            self.assertIn(key, payload, f"evaluate-model: missing top-level key '{key}'")
        for key in schema["required_metric_keys"]:
            self.assertIn(key, payload["metrics"], f"evaluate-model: missing metric '{key}'")
        for key in schema["required_confusion_matrix_keys"]:
            self.assertIn(key, payload["confusion_matrix"], f"evaluate-model: missing confusion_matrix key '{key}'")
        for key in schema["required_split_keys"]:
            self.assertIn(key, payload["split"], f"evaluate-model: missing split key '{key}'")

    def test_run_data_pipeline_output_satisfies_contract(self) -> None:
        schema = self._load_schema("run-data-pipeline.schema.json")
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = self.write_analytics_csv(root)
            output = root / "pipeline_out"
            result = self.run_cli(
                "--json", "run", "data-scientist-analyst", "run-data-pipeline",
                "--source", str(source),
                "--output", str(output),
            )
        self.assertEqual(result.returncode, 0, result.stderr)
        payload = self.parse_payload(result)
        for key in schema["required_top_level_keys"]:
            self.assertIn(key, payload, f"run-data-pipeline: missing top-level key '{key}'")
        for key in schema["required_pipeline_keys"]:
            self.assertIn(key, payload["pipeline"], f"run-data-pipeline: missing pipeline key '{key}'")
        self.assertTrue(payload["pipeline"]["cache_key"], "cache_key must be non-empty")
        self.assertIn(payload["pipeline"]["status"], {"success", "warning"})

    def test_system_md_exists_and_is_in_default_context(self) -> None:
        """P0.1 — knowledge/system.md exists and agent.yaml references it."""
        agent_root = ROOT / "agents" / "data-scientist-analyst"
        system_md = agent_root / "knowledge" / "system.md"
        self.assertTrue(system_md.exists(), "knowledge/system.md must exist")
        content = system_md.read_text(encoding="utf-8")
        self.assertGreater(len(content.strip()), 200, "system.md must not be a placeholder")
        import yaml
        agent_yaml = yaml.safe_load((agent_root / "agent.yaml").read_text(encoding="utf-8"))
        self.assertIn(
            "knowledge/system.md",
            agent_yaml["default_context"],
            "knowledge/system.md must be in default_context",
        )
        self.assertEqual(
            agent_yaml["default_context"][0],
            "knowledge/system.md",
            "knowledge/system.md must be the FIRST item in default_context",
        )

    def test_health_checklist_is_in_default_context(self) -> None:
        """P0.3 — health-checklist.md promoted to default_context."""
        agent_root = ROOT / "agents" / "data-scientist-analyst"
        import yaml
        agent_yaml = yaml.safe_load((agent_root / "agent.yaml").read_text(encoding="utf-8"))
        self.assertIn(
            "knowledge/health-checklist.md",
            agent_yaml["default_context"],
            "knowledge/health-checklist.md must be in default_context",
        )

    def test_all_prompts_are_non_placeholder(self) -> None:
        """P0.2 — no prompt file may be a placeholder (< 200 chars)."""
        prompts_dir = ROOT / "agents" / "data-scientist-analyst" / "knowledge" / "prompts"
        for prompt_file in sorted(prompts_dir.glob("*.md")):
            content = prompt_file.read_text(encoding="utf-8").strip()
            self.assertGreater(
                len(content),
                200,
                f"{prompt_file.name} is too short (placeholder); must follow Section 5 format",
            )

    def test_policies_yaml_exists_and_covers_required_fields(self) -> None:
        """P1.1 — knowledge/policies.yaml exists with write_policy and pii_masking."""
        import yaml
        policies_path = ROOT / "agents" / "data-scientist-analyst" / "knowledge" / "policies.yaml"
        self.assertTrue(policies_path.exists(), "knowledge/policies.yaml must exist")
        policies = yaml.safe_load(policies_path.read_text(encoding="utf-8"))
        self.assertIn("write_policy", policies)
        self.assertIn("pii_masking", policies)
        self.assertIn("decision_rules", policies)
        self.assertEqual(policies["write_policy"]["source_mutation"], "unsupported")
        masked = policies["pii_masking"]["mask_if_sensitive"]
        for field in ["cpf", "cnpj", "email", "telefone", "nome"]:
            self.assertIn(field, masked, f"pii_masking must cover '{field}'")

    def test_schema_contracts_referenced_in_capability_yaml(self) -> None:
        """P1.2 — the 4 schema contracts are referenced by their capability.yaml."""
        import yaml
        agent_root = ROOT / "agents" / "data-scientist-analyst"
        contract_map = {
            "profile-dataset": "profile-dataset.schema.json",
            "evaluate-model": "evaluate-model.schema.json",
            "run-data-pipeline": "run-data-pipeline.schema.json",
            "analyze-sql-source": "analyze-sql-source.schema.json",
        }
        for cap_id, schema_file in contract_map.items():
            cap_yaml_path = agent_root / "capabilities" / cap_id / "capability.yaml"
            cap_yaml = yaml.safe_load(cap_yaml_path.read_text(encoding="utf-8"))
            output_template = cap_yaml.get("entrypoint", {}).get("output_template", "")
            self.assertIn(
                schema_file,
                output_template or "",
                f"{cap_id}/capability.yaml must reference {schema_file} in output_template",
            )

    def test_templates_are_not_placeholders(self) -> None:
        """P1.3 — output templates are filled, not one-line placeholders."""
        agent_root = ROOT / "agents" / "data-scientist-analyst"
        for name in ["data-report.md", "reconciliation-report.md"]:
            content = (agent_root / "templates" / name).read_text(encoding="utf-8").strip()
            self.assertGreater(
                len(content),
                200,
                f"templates/{name} is a placeholder; must document the output contract",
            )


if __name__ == "__main__":
    unittest.main()
